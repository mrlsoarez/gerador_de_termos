from CLASSES.Documento import Documento
from docx import Document
import locale 

class Relatorio(Documento):
    
    
    def __init__(self, contratado, liquidacao, valor, data, modelo, protocolo, endereco):
        super().__init__(contratado)
        self.liquidacao = liquidacao
        self.valor = valor
        self.data = data 
        self.modelo = modelo
        self.protocolo = protocolo
        self.endereco = endereco
    
    def criar_arquivo(self, termos):
        
        doc = Document(self.endereco)
        
        def formatar_data(self, objeto):
            return objeto.date().strftime("%d/%m/%Y")
        
        def converter_currency(self, valor):
            locale.setlocale(locale.LC_ALL, "pt_BR.UTF-8")
            return locale.currency(float(valor), grouping =True)
                    
        def adicionar_protocolo(self):
            substituir_protocolo = doc.paragraphs[1]
            substituir_protocolo.text = ""
            Documento.criar_texto(substituir_protocolo, f"PROTOCOLO DE RECEBIMENTO - NÚMERO {self.protocolo}", negrito = True)
               
        def criar_tabela(self, termos):
            
            for i in range(len(termos)):
                
                tabela = doc.tables[0]   
                nova_linha = tabela.add_row()
                                
                coluna_um = nova_linha.cells[0].paragraphs[0]
                coluna_dois = nova_linha.cells[1].paragraphs[0]
                coluna_tres = nova_linha.cells[2].paragraphs[0]
                coluna_quatro = nova_linha.cells[3].paragraphs[0]
                                    
                Documento.criar_texto(coluna_um, termos[i].contratado,  px = 8, negrito = True, fonte = "Arial")
                Documento.criar_texto(coluna_dois, termos[i].liquidacao,  px = 8, negrito = True, fonte = "Arial")
                Documento.criar_texto(coluna_tres, formatar_data(self, termos[i].data),  px = 8, negrito = True, fonte = "Arial")
                Documento.criar_texto(coluna_quatro, converter_currency(self, termos[i].valor),  px = 8, negrito = True, fonte = "Arial")
                                
        adicionar_protocolo(self) 
        criar_tabela(self, termos)
        
        doc.save(self.endereco)