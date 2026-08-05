from CLASSES.Documento import Documento
from docx import Document
from docx2pdf import convert

class Termo(Documento):
    
    def __init__(self, contratado, endereco, ordem, contrato, objeto, af, mensagem, gestor, tipo, modelo):
        super().__init__(contratado)
        self.ordem = ordem
        self.endereco = endereco
        self.contrato = contrato 
        self.objeto = objeto 
        self.af = af
        self.mensagem = mensagem
        self.gestor = gestor
        self.tipo = tipo
        self.modelo = modelo 

    def criar_arquivo(self):
        
        doc = Document(self.modelo)
        
        def definir_tabela(self):

            if self.tipo.lower() == "ata":
                campo = "ATA N°"
            else:
                campo = "CONTRATO N°" 

            dados = [{"celula": (1, 0), "conteudo": campo, "negrito": True}, 
                     {"celula": (1, 1), "conteudo": self.contrato, "negrito": False }, 
                     {"celula": (2, 1), "conteudo": self.contratado, "negrito": False }, 
                     {"celula": (3, 1), "conteudo": self.objeto, "negrito": False }, 
                     {"celula": (4, 1), "conteudo": self.af, "negrito": False}, 
                     {"celula": (6, 1), "conteudo": self.mensagem, "negrito": False }, 
                    ]
            
            doc.tables[0] = Documento.modificar_tabela(doc.tables[0], dados)
            
        def definir_data(self):
            data_texto = "BATAGUASSU/MS, " + Documento.encontrar_data_de_hoje_em_extenso()
            Documento.criar_texto(doc.add_paragraph(), data_texto, negrito = True, posicionamento = "Direita")
        
        def adicionar_espaco(self, quant):
            for i in range(quant):
                doc.add_paragraph("")
        
        def definir_gestor(self):
            Documento.adicionar_linha_de_assinatura(doc.add_paragraph(), self.gestor)
            
                
        definir_tabela(self)
        definir_data(self)
        adicionar_espaco(self, 3)
        definir_gestor(self)
        
        doc.save(self.endereco)
    
    def salvar_pdf(self):
        docx = self.endereco
        pdf = docx.replace("WORD", "PDF")
        pdf = pdf.replace(".docx", ".pdf")
        print(f"Convertendo termo para PDF.... *.✧*.✧.*.✧*.✧*.✧*.✧.,*.✧*. {self.contratado} - AF {self.af}")
        try:
            convert(docx, pdf)
        except:
            pass 