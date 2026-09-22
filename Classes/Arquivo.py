

from openpyxl import load_workbook

#from MODULES.GerenciarArquivos import GerenciarArquivos
from Services.Calendario import EncontrarData


import shutil
import locale

from docx import Document
from docx2pdf import convert
from copy import deepcopy

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from datetime import date

class Arquivo: 
    
    def __init__(self):
        pass 

    def set_arq_nome(self, dado):
        self.arq = dado 
        
    def copiar_arquivo(self, antigo, novo):
        shutil.copy(antigo, novo)
    
    def mudar_fonte(texto, name_font):
        texto.font.name = name_font
    
    def mudar_tamanho(texto, px): 
        texto.font.size = Pt(px)

    @staticmethod
    def criar_texto(paragrafo, texto, negrito = False, posicionamento = None, px = None, fonte = None):

        def deixar_negrito(run): 
           run.bold = True 

        def alinhar_texto(paragrafo, alinhado):
            if (alinhado == "Centro"): 
                paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif (alinhado == "Direita"):
                paragrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT 

        run = paragrafo.add_run(str(texto))
        if (negrito): deixar_negrito(run)
        if (posicionamento != None): alinhar_texto(paragrafo, posicionamento)
        if (px != None): Arquivo.mudar_tamanho(run, px)
        if (fonte != None): Arquivo.mudar_fonte(run, fonte)

    @staticmethod
    def encontrar_data_de_hoje_em_extenso():  

        dia = EncontrarData("dia")
        mes = EncontrarData("mes", True)
        ano = EncontrarData("ano")
        
        return f"{dia} de {mes} de {ano}"

    @staticmethod
    def adicionar_linha_de_assinatura(paragrafo, assinador): 
        Arquivo.criar_texto(paragrafo, f"________________________________________\n{assinador}", negrito = True, posicionamento = "Centro", fonte = "Cambria")

    @staticmethod
    def modificar_tabela(tabela, dados, fonte = "Cambria"): 

        for i in range(len(dados)):

            dado = dados[i]
                
            celula = tabela.cell(dado["celula"][0], dado["celula"][1])
            paragrafo = celula.paragraphs[0]  

            Arquivo.criar_texto(paragrafo, dado["conteudo"], dado["negrito"], fonte, px = 10)

        return tabela
    
    @staticmethod
    def salvarPDF(endereco):
        docx = endereco
        pdf = docx.replace(".docx", ".pdf")
        print(docx, pdf)
        try:
            convert(docx, pdf)
        except Exception as e:
            print(e)
            pass 
