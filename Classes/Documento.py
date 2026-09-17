from docx import Document  
import os 
import locale 


from Classes.Arquivo import Arquivo

class Documento():

    mapa = {
        "contratado": "B4",
        "n_contrato": "E4",
        "objeto": "F4",
        "numero_empenho": "A8",
        "numero_liquidacao": "A12",
        "data_liquidacao": "B12",
        "valor_bruto_liquidacao": "C12",
        "tipo_nota": "D16",
        "numero_af": "A20",
        "tipo": "B5"
    }

    def __init__(self, ordem, contratado, contrato, objeto, af, tipo_nota, liq, data, valor, tipo):
        # Info Termo
        self.ordem = ordem 
        self.contratado = contratado
        self.contrato = contrato 
        self.objeto = objeto 
        self.af = af
        self.tipo_nota = tipo_nota

        # Info Relatorio
        self.liq = liq 
        self.data = data 
        self.valor = valor

        self.tipo = tipo

class Termo(Documento): 

    def __init__(self, termo, modelo, endereco):
        self.termo = termo 
        self.modelo = modelo 
        self.endereco = endereco 
        pass 
    
    def setEndereco(self, endereco): 
        self.endereco = endereco
    
    def setQuantidadeTermos(self, quant):
        self.quantidade = quant + 1

    def verificarSeExiste(self): 
        info = self.termo     
        nome_arquivo = rf"{self.ordem}. {info.contratado} - AF {info.af[:4]}.docx"
        self.setEndereco(rf"{self.endereco}\{nome_arquivo}")
        return os.path.exists(self.endereco)
        
    def setOrdem(self, ordem):
        self.ordem = ordem
                
    def setGestor(self): 
        if (self.termo.tipo.lower() == "contrato"):
            self.gestor = "RONALDO DE SOUZA MARCÍLIO\nGESTOR DE CONTRATOS"   
        else:
            self.gestor = "MURILO SOARES DE OLIVEIRA\nGESTOR DE ATAS"
        pass 
    
    def setMensagem(self):
        if (self.termo.tipo.lower() == "locação"): 
            self.mensagem = f"Por este instrumento, em caráter DEFINITIVO, atestamos que a locação acima identificada atende às exigências contratuais."
            return   
        self.mensagem = f"Por este instrumento, em caráter DEFINITIVO, atestamos que os {self.termo.tipo_nota.lower()} acima identificados atendem às exigências contratuais."
        pass

    def criarArquivo(self):
        
        doc = Document(self.modelo)
        info = self.termo 
        
        def definir_tabela(self):

            if info.tipo.lower() == "ata":
                campo = "ATA N°"
            else:
                campo = "CONTRATO N°" 

            dados = [{"celula": (1, 0), "conteudo": campo, "negrito": True}, 
                     {"celula": (1, 1), "conteudo": info.contrato, "negrito": False }, 
                     {"celula": (2, 1), "conteudo": info.contratado, "negrito": False }, 
                     {"celula": (3, 1), "conteudo": info.objeto, "negrito": False }, 
                     {"celula": (4, 1), "conteudo": info.af, "negrito": False}, 
                     {"celula": (6, 1), "conteudo": self.mensagem, "negrito": False }, 
                    ]
            
            doc.tables[0] = Arquivo.modificar_tabela(doc.tables[0], dados)
            
        def definir_data(self):
            data_texto = "BATAGUASSU/MS, " + Arquivo.encontrar_data_de_hoje_em_extenso()
            Arquivo.criar_texto(doc.add_paragraph(), data_texto, negrito = True, posicionamento = "Direita")
        
        def adicionar_espaco(self, quant):
            for i in range(quant):
                doc.add_paragraph("")
        
        def definir_gestor(self):
            Arquivo.adicionar_linha_de_assinatura(doc.add_paragraph(), self.gestor)
        
        self.setGestor()
        self.setMensagem() 
           
        definir_tabela(self)
        definir_data(self)
        adicionar_espaco(self, 3)
        definir_gestor(self)
        
        try: 
            doc.save(self.endereco)
        except Exception as e:
            print(e)
        else: 
            print(rf"Documento salvo: {self.endereco}, convertendo para PDF...")
        
        Arquivo.salvarPDF(self.endereco)

class Relatorio(Documento):

    def __init__(self, relatorio, n_protocolo, modelo, endereco):
        self.relatorio = relatorio 
        self.n_protocolo = n_protocolo
        self.modelo = modelo 
        self.endereco = endereco
        pass   
    
    def setEndereco(self, endereco):
        self.endereco = endereco 
        
    def criarArquivo(self, termos):
        
        doc = Document(self.modelo)
        
        def formatar_data(self, objeto):
            return objeto.date().strftime("%d/%m/%Y")
        
        def converter_currency(self, valor):
            locale.setlocale(locale.LC_ALL, "pt_BR.UTF-8")
            return locale.currency(float(valor), grouping =True)
                    
        def adicionar_protocolo(self):
            substituir_protocolo = doc.paragraphs[1]
            substituir_protocolo.text = ""
            Arquivo.criar_texto(substituir_protocolo, f"PROTOCOLO DE RECEBIMENTO - NÚMERO {self.n_protocolo}", negrito = True)
               
        def criar_tabela(self, termos):
            
            for i in range(len(termos)):
                
                tabela = doc.tables[0]   
                nova_linha = tabela.add_row()
                                
                coluna_um = nova_linha.cells[0].paragraphs[0]
                coluna_dois = nova_linha.cells[1].paragraphs[0]
                coluna_tres = nova_linha.cells[2].paragraphs[0]
                coluna_quatro = nova_linha.cells[3].paragraphs[0]
                Arquivo.criar_texto(coluna_um, termos[i].contratado,  px = 8, negrito = True, fonte = "Arial")
                Arquivo.criar_texto(coluna_dois, termos[i].liq,  px = 8, negrito = True, fonte = "Arial")
                Arquivo.criar_texto(coluna_tres, formatar_data(self, termos[i].data),  px = 8, negrito = True, fonte = "Arial")
                Arquivo.criar_texto(coluna_quatro, converter_currency(self, termos[i].valor),  px = 8, negrito = True, fonte = "Arial")
                                
        adicionar_protocolo(self) 
        criar_tabela(self, termos)
        doc.save(rf"{self.endereco}\PROTOCOLO DE RECEBIMENTO - N° {self.n_protocolo}.docx")