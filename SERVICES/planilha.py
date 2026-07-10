
import win32com.client

from SERVICES.ARQUIVO import PROCESSAR_ARQUIVO
import pandas as pd
import os

from datetime import datetime
from openpyxl import load_workbook
import time

# Funções referentes a planilha

def CONVERTER_PARA_XLSX(arquivo):
        arquivo_xls = arquivo
        arquivo_xlsx = arquivo_xls + "x"
        try:
            df = pd.read_excel(arquivo_xls, engine="xlrd")
            df.to_excel(arquivo_xlsx, index=False)
            os.remove(arquivo_xls)
        except:
            pass 
        finally: 
            return arquivo_xlsx
            
        
def ATUALIZAR_PLANILHA_COM_DADOS_EXTERNOS(localizacao_planilha, nome_sheet, dados):
    
    ANALISE_FISCAL = load_workbook(localizacao_planilha)
    
    def limpar_planilha(sheet):
        if sheet.max_row > 0:
            sheet.delete_rows(1, sheet.max_row)
    
    SHEET = ANALISE_FISCAL[nome_sheet]
    limpar_planilha(SHEET)

    
    HEAD = []
    for keys in dados[0]:
        HEAD.append(keys)
    SHEET.append(HEAD)

    for linha in dados: 
        SHEET.append([linha.get(coluna, "") for coluna in HEAD])
        
   
    ANALISE_FISCAL.save(localizacao_planilha)
    """


    def atualizar_planilha(sheet_origem, sheet_destino):
        origem = sheet_origem.active
        destino = ANALISE_FISCAL[sheet_destino]
       
        for linha_idx, linha in enumerate(origem.iter_rows(values_only=True), start=1):
            for coluna_idx, valor in enumerate(linha, start=1):
                destino.cell(
                    row=linha_idx,
                    column=coluna_idx,
                    value=valor
                )
    
    try:
        converter_para_xlsx(rf"{pasta}\Portal Transparencia Despesas Gerais - Exercício 2026.xls")
        converter_para_xlsx(rf"{pasta}\Portal Transp. Despesas Liquidadas.xls")
    except: 
        pass

    ANALISE_FISCAL = load_workbook(localizacao_planilha)
    planilha_empenhos = load_workbook(rf"{pasta}\Portal Transparencia Despesas Gerais - Exercício 2026.xlsx")
    planilha_liquidacao = load_workbook(rf"{pasta}\Portal Transp. Despesas Liquidadas.xlsx")

    try:
        atualizar_planilha(planilha_empenhos, "Empenhos")
        atualizar_planilha(planilha_liquidacao, "Liquidacoes")
    except:
        print("Não foi possível completar a substituição dos dados na planilha.")
    else:
        ANALISE_FISCAL.save(localizacao_planilha)

    """
   
    
def INICIAR_PLANILHA(localizacao_planilha, gerenciador_arquivos):

    excel = win32com.client.Dispatch("Excel.Application")
    excel.Visible = True

    wb = excel.Workbooks.Open(localizacao_planilha)
    
    modified = os.path.getmtime(localizacao_planilha)
    modified = datetime.fromtimestamp(modified)
    try:
        while True:
            time.sleep(1)
            last_modified = os.path.getmtime(localizacao_planilha)
            last_modified = datetime.fromtimestamp(last_modified)
            if (last_modified > modified):
                # aqui começa a criar o termo, op assincrona?
                PROCESSAR_ARQUIVO(localizacao_planilha, gerenciador_arquivos)
                modified = os.path.getmtime(localizacao_planilha)
                modified = datetime.fromtimestamp(modified)
                gerenciador_arquivos.entrar_na_pasta("..")
                
    except KeyboardInterrupt:
        print("Encerrando...")
    finally:
        wb.Close(SaveChanges=True)  
        excel.Quit()               
