
from openpyxl import load_workbook

from MODULES.GerenciarArquivos import GerenciarArquivos
from MODULES.EncontrarData import EncontrarData

from ENV.environment import pegar_modelos

import shutil
import locale

from docx import Document
from docx2pdf import convert
from copy import deepcopy

from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import date
from docx.shared import Pt


class Documento: 
    
    def __init__(self, contratado):
        self.contratado = contratado
        
    def set_af(self, dado):
        self.af = dado
    
    def set_arq_nome(self, dado):
        self.arq = dado 
        
    def copiar_arquivo(self, antigo, novo):
        shutil.copy(antigo, novo)
    

    
    def mudar_fonte(texto, name_font):
        texto.font.name = name_font
    
    def mudar_tamanho(texto, px): 
        texto.font.size = Pt(px)

    @staticmethod
    def criar_texto(paragrafo, texto, negrito = False, posicionamento = None, px = None, fonte = None):

        def deixar_negrito(run): 
           run.bold = True 

        def alinhar_texto(paragrafo, alinhado):
            if (alinhado == "Centro"): 
                paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif (alinhado == "Direita"):
                paragrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT 

        run = paragrafo.add_run(str(texto))
        if (negrito): deixar_negrito(run)
        if (posicionamento != None): alinhar_texto(paragrafo, posicionamento)
        if (px != None): Documento.mudar_tamanho(run, px)
        if (fonte != None): Documento.mudar_fonte(run, fonte)

    @staticmethod
    def encontrar_data_de_hoje_em_extenso():  

        dia = EncontrarData("dia")
        mes = EncontrarData("mes", True)
        ano = EncontrarData("ano")
        
        return f"{dia} de {mes} de {ano}"

    @staticmethod
    def adicionar_linha_de_assinatura(paragrafo, assinador): 
        Documento.criar_texto(paragrafo, f"________________________________________\n{assinador}", negrito = True, posicionamento = "Centro", fonte = "Cambria")
    
    @staticmethod
    def modificar_tabela(tabela, dados, fonte = "Cambria"): 

        for i in range(len(dados)):

            dado = dados[i]
                
            celula = tabela.cell(dado["celula"][0], dado["celula"][1])
            paragrafo = celula.paragraphs[0]  

            Documento.criar_texto(paragrafo, dado["conteudo"], dado["negrito"], fonte, px = 10)

        return tabela
    
    @staticmethod
    def salvar_pdf(self):
        pdf = self.endereco.replace("WORD", "PDF")
        pdf = pdf.replace(".docx", ".pdf")
        #convert(docx, pdf)
        return pdf
                   
class Termo(Documento):
    
    def __init__(self, contratado, endereco, ordem, contrato, objeto, af, mensagem, gestor, tipo, modelo):
        super().__init__(contratado)
        self.ordem = ordem
        self.endereco = endereco
        self.contrato = contrato 
        self.objeto = objeto 
        self.af = af
        self.mensagem = mensagem
        self.gestor = gestor
        self.tipo = tipo
        self.modelo = modelo 

    def criar_arquivo(self):
        
        doc = Document(self.modelo)
        
        def definir_tabela(self):

            if self.tipo.lower() == "ata":
                campo = "ATA N°"
            else:
                campo = "CONTRATO N°" 

            dados = [{"celula": (1, 0), "conteudo": campo, "negrito": True}, 
                     {"celula": (1, 1), "conteudo": self.contrato, "negrito": False }, 
                     {"celula": (2, 1), "conteudo": self.contratado, "negrito": False }, 
                     {"celula": (3, 1), "conteudo": self.objeto, "negrito": False }, 
                     {"celula": (4, 1), "conteudo": self.af, "negrito": False}, 
                     {"celula": (6, 1), "conteudo": self.mensagem, "negrito": False }, 
                    ]
            
            doc.tables[0] = Documento.modificar_tabela(doc.tables[0], dados)
            
        def definir_data(self):
            data_texto = "BATAGUASSU/MS, " + Documento.encontrar_data_de_hoje_em_extenso()
            Documento.criar_texto(doc.add_paragraph(), data_texto, negrito = True, posicionamento = "Direita")
        
        def adicionar_espaco(self, quant):
            for i in range(quant):
                doc.add_paragraph("")
        
        def definir_gestor(self):
            Documento.adicionar_linha_de_assinatura(doc.add_paragraph(), self.gestor)
            
                
        definir_tabela(self)
        definir_data(self)
        adicionar_espaco(self, 3)
        definir_gestor(self)
        
        doc.save(self.endereco)
    
    def salvar_pdf(self):
        docx = self.endereco
        pdf = docx.replace("WORD", "PDF")
        pdf = pdf.replace(".docx", ".pdf")
        print(f"Convertendo termo para PDF.... *.✧*.✧.*.✧*.✧*.✧*.✧.,*.✧*. {self.contratado} - AF {self.af}")
        try:
            convert(docx, pdf)
        except:
            pass 
        
class Relatorio(Documento):
    
    
    def __init__(self, contratado, liquidacao, valor, data, modelo, protocolo, endereco):
        super().__init__(contratado)
        self.liquidacao = liquidacao
        self.valor = valor
        self.data = data 
        self.modelo = modelo
        self.protocolo = protocolo
        self.endereco = endereco
    
    def criar_arquivo(self, termos):
        
        doc = Document(self.endereco)
        
        def formatar_data(self, objeto):
            return objeto.date().strftime("%d/%m/%Y")
        
        def converter_currency(self, valor):
            locale.setlocale(locale.LC_ALL, "pt_BR.UTF-8")
            return locale.currency(float(valor), grouping =True)
                    
        def adicionar_protocolo(self):
            substituir_protocolo = doc.paragraphs[1]
            substituir_protocolo.text = ""
            Documento.criar_texto(substituir_protocolo, f"PROTOCOLO DE RECEBIMENTO - NÚMERO {self.protocolo}", negrito = True)
               
        def criar_tabela(self, termos):
            
            for i in range(len(termos)):
                
                tabela = doc.tables[0]   
                nova_linha = tabela.add_row()
                                
                coluna_um = nova_linha.cells[0].paragraphs[0]
                coluna_dois = nova_linha.cells[1].paragraphs[0]
                coluna_tres = nova_linha.cells[2].paragraphs[0]
                coluna_quatro = nova_linha.cells[3].paragraphs[0]
                                    
                Documento.criar_texto(coluna_um, termos[i].contratado,  px = 8, negrito = True, fonte = "Arial")
                Documento.criar_texto(coluna_dois, termos[i].liquidacao,  px = 8, negrito = True, fonte = "Arial")
                Documento.criar_texto(coluna_tres, formatar_data(self, termos[i].data),  px = 8, negrito = True, fonte = "Arial")
                Documento.criar_texto(coluna_quatro, converter_currency(self, termos[i].valor),  px = 8, negrito = True, fonte = "Arial")
                                
        adicionar_protocolo(self) 
        criar_tabela(self, termos)
        
        doc.save(self.endereco)
 
        
class Portaria(Documento):
        def __init__(self, numero_portaria, codigo_tce, fornecedores, objeto, valor_total, modelo, endereco):
            self.numero_portaria = numero_portaria
            self.codigo_tce = codigo_tce
            self.fornecedores = fornecedores
            self.objeto = objeto
            self.valor_total = valor_total
            self.fiscais = [] 
            self.modelo = modelo 
            self.endereco = endereco
            
        def criar_arquivo(self):
            doc = Document(self.modelo)
            
            def adicionar_titulo():
                doc.paragraphs[1].text = ""
                titulo = doc.paragraphs[1]
                texto = f"Portaria N° {self.numero_portaria} - {Documento.encontrar_data_de_hoje_em_extenso()}"
                Documento.criar_texto(titulo, texto, negrito = True, posicionamento="Centro", fonte = "Arial")
                
            def adicionar_codigo_tce():
                doc.paragraphs[4].text = ""
                tce = doc.paragraphs[4]
                texto = f"Código Registro TCE: [{self.codigo_tce}]"
                Documento.criar_texto(tce, texto, negrito = True, posicionamento="Centro", fonte = "Arial")
        
            def criar_tabela_fiscais():
                    
                    def capitalizar(texto):
                        excecoes = {"de", "e"}

                        return " ".join(
                            palavra if palavra.lower() in excecoes
                            else palavra.capitalize()
                            for palavra in texto.lower().split()
                        )
                        
                    def criar_nova_linha(dados):
                        nova_linha = tabela.add_row()
                        
                        coluna_um = nova_linha.cells[0].paragraphs[0]
                        coluna_dois = nova_linha.cells[1].paragraphs[0]
                        coluna_tres = nova_linha.cells[2].paragraphs[0]
                        coluna_quatro = nova_linha.cells[3].paragraphs[0]
                                                
                        Documento.criar_texto(coluna_um, dados[0],  px = 11, fonte = "Arial")
                        Documento.criar_texto(coluna_dois, dados[1],  px = 11,  fonte = "Arial")
                        Documento.criar_texto(coluna_tres, dados[2],  px = 11,  fonte = "Arial")
                        Documento.criar_texto(coluna_quatro, dados[3],  px = 11, fonte = "Arial")
                        
                    fiscais = self.fiscais  
                                       
                    tabela = doc.tables[0]   
                    
                    for index in range(len(fiscais)):
                        
                        criar_nova_linha(["", "FISCAL", "SUPLENTE", "GESTOR"])
                        criar_nova_linha(["NOME DO SERVIDOR", capitalizar(fiscais[index]["principal"].nome), capitalizar(fiscais[index]["suplente"].nome), "Murilo Soares de Oliveira"])
                        criar_nova_linha(["CARGO", capitalizar(fiscais[index]["principal"].cargo), capitalizar(fiscais[index]["suplente"].cargo), "Assistente de Administração"])
                        criar_nova_linha(["MATRÍCULA", fiscais[index]["principal"].matricula, fiscais[index]["suplente"].matricula, "117810"])
                        criar_nova_linha(["VÍNCULO", capitalizar(fiscais[index]["principal"].vinculo), capitalizar(fiscais[index]["suplente"].vinculo), "Efetivo"])
                        criar_nova_linha(["SECRETARIA", capitalizar(fiscais[index]["principal"].secretaria), capitalizar(fiscais[index]["suplente"].secretaria), "Administração e Finanças"])

                        if (index != len(fiscais) - 1): 
                            nova_linha = tabela.add_row()
            
            def alterar_tabela_atas():
                
                tabela = doc.tables[1] 
                texto = ""
                fornecedores = self.fornecedores 
                
                for i in range(len(fornecedores)):
                    texto += f"ATA N° {fornecedores[i]['ata']} - {fornecedores[i]['fornecedor']}, CNPJ N° {fornecedores[i]['cnpj']}\n"
                
                dados = [
                    {"celula": (0, 1), "conteudo": texto, "negrito": True },
                    {"celula": (1, 1), "conteudo": self.objeto, "negrito": True },
                    {"celula": (2, 1), "conteudo": "12 meses.", "negrito": True },
                    {"celula": (3, 1), "conteudo": f"R$ {self.valor_total}", "negrito": True },
               
                ]  
                            
                doc.tables[1] = Documento.modificar_tabela(doc.tables[1], dados, "Arial")
                pass   
            
            def inserir_data_rodape():
                index = None
                for i in range(len(doc.paragraphs)):
                    if ("Gabinete da Prefeita Municipal de Bataguassu-MS, em [DATA]" in doc.paragraphs[i].text):
                        index = i 
                        
                doc.paragraphs[index].text = ""
                paragrafo = doc.paragraphs[index] 
                texto = f"Gabinete da Prefeita Municipal de Bataguassu-MS, em {Documento.encontrar_data_de_hoje_em_extenso()}"
                Documento.criar_texto(paragrafo, texto, negrito = True, fonte = "Arial")
                pass 
            
            adicionar_titulo()
            adicionar_codigo_tce()
            criar_tabela_fiscais()
            alterar_tabela_atas()
            inserir_data_rodape() 
            
            doc.save(rf"{self.endereco}\Portaria N° xx.2026 - Objeto.docx")
            
               
class Fiscal: 
    def __init__(self, nome, matricula, cargo, vinculo, secretaria):
        self.nome = nome 
        self.matricula = matricula 
        self.cargo = cargo 
        self.vinculo = vinculo
        self.secretaria = secretaria
        
    def get_fiscal(self):
        return self.fiscal
                