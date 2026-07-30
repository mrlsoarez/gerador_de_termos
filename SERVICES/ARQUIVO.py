from openpyxl import load_workbook

from MODULES.GerenciarArquivos import GerenciarArquivos
from MODULES.EncontrarData import EncontrarData

from ENV.environment import pegar_modelos
from SERVICES.Classes.Documento import Documento, Termo, Relatorio, Portaria, Fiscal
from SERVICES.Classes.Verificador import Verificador 

import os
import shutil
from docx import Document
from docx2pdf import convert

from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import date
from docx.shared import Pt

protocolo_existe = False 

modelo_relatorio = pegar_modelos("relatorio")

# she was so happy to look a mess
# affs
def PROCESSAR_ARQUIVO(sheet_planilha, gerenciador, param):
    
    OPTION = param["op"]
    
    def gerar_termos(sheet, ordem):
                
        verificador = Verificador("", sheet)   

        def verificar_informacoes_iniciais():
            verificacao = verificador.checar_campos_planilha()
            
            if (verificacao["resultado"]):
                verificador.set_info()
                if (not gerenciador.verificarArquivo(verificador.arq)):
                    doc = Documento(verificador.contratado)
                    doc.set_af(verificador.af)
                    doc.set_arq_nome(verificador.arq)
                    return doc
                return False
            else: 
                print(verificacao["mensagem"])
                return False 
                
        def colher_informacoes_termo(doc):
            
            def customizar_mensagem(tipo):
                if (tipo == "Locação"): return "Por este instrumento, em caráter DEFINITIVO, atestamos que a locação acima identificada atende às exigências contratuais."
                return f"Por este instrumento, em caráter DEFINITIVO, atestamos que os {tipo.lower()} acima identificados atendem às exigências contratuais."
            
            def definir_gestor(tipo):
                if (tipo == "contrato"):
                    return "RONALDO DE SOUZA MARCÍLIO\nGESTOR DE CONTRATOS"   
                else:
                    return "MURILO SOARES DE OLIVEIRA\nGESTOR DE ATAS"
            
            mapa = verificador.mapa 
            
            mensagem = customizar_mensagem(str(sheet[mapa['tipo_nota']].value))
            endereco = rf"{gerenciador.pasta_atual}\{doc.arq}"
            gestor = definir_gestor(gerenciador.tipo_arquivo)
            contrato = sheet[mapa['n_contrato']].value
            modelo_termo = pegar_modelos("termo")
            objeto = sheet[mapa['objeto']].value 
            tipo = gerenciador.tipo_arquivo
            contratado = doc.contratado
            numero_af = doc.af 
            
            return Termo(contratado, endereco, ordem, contrato, objeto, numero_af, mensagem, gestor, tipo, modelo_termo)
    
        doc_verificacao = verificar_informacoes_iniciais()
        
        if (doc_verificacao != False):
            termo = colher_informacoes_termo(doc_verificacao)
            termo.criar_arquivo()    
            termo.salvar_pdf()     
    
    def buscar_informacoes_protocolo(sheet):
        
        verificador = Verificador("", sheet)   
        def verificar_informacoes_iniciais():
            verificacao = verificador.checar_campos_planilha()
            if (verificacao["resultado"]):
                return True
            
            print(verificacao["mensagem"])
        
        def colher_informacao_relatorio():
                    
            mapa = verificador.mapa         
                        
            contratado = sheet[mapa['contratado']].value
            liquidacao = sheet[mapa['numero_liquidacao']].value
            data = sheet[mapa['data_liquidacao']].value 
            valor = sheet[mapa['valor_bruto_liquidacao']].value 
            modelo_termo = pegar_modelos("protocolo")
            endereco = rf"{gerenciador.pasta_atual}\Protocolo N° {gerenciador.numero_protocolo} - Tesouraria.docx"
            
            return Relatorio(contratado, liquidacao, valor, data, modelo_termo, gerenciador.numero_protocolo, endereco)
    
        verificacao = verificar_informacoes_iniciais()
        
        if (verificacao):
            print(verificacao)
            relatorio = colher_informacao_relatorio()
            param["ARR"].append(relatorio)
        pass         
    
    def gerar_portarias(sheet):
        
        def coletar_informacoes_portarias():
            
            portaria = None
            working_sheet = sheet["Portarias"]
            find_fiscais = False 
            
            fornecedores = []
            fiscais = []
            
            objeto = working_sheet["E2"].value
            cod_tce = working_sheet["F2"].value
            valor = working_sheet["G2"].value
            modelo = pegar_modelos("portaria")
            
            for i in range(2, working_sheet.max_row):
                if (working_sheet["A" + str(i)].value == "Portaria"):
                    find_fiscais = True
                    portaria = Portaria("XX/26", cod_tce, fornecedores, objeto, valor, modelo, gerenciador.pasta_base)
                    continue
                
                if (find_fiscais):
                    if (working_sheet["A" + str(i)].value != None):
                        
                        principal = working_sheet["B" + str(i)].value
                        matricula_p = working_sheet["C" + str(i)].value
                        cargo_p = working_sheet["D" + str(i)].value
                        secretaria_p = working_sheet["E" + str(i)].value
                        vinculo_p = working_sheet["F" + str(i)].value
                        
                        
                        suplente = working_sheet["G" + str(i)].value
                        matricula_s = working_sheet["H" + str(i)].value
                        cargo_s = working_sheet["I" + str(i)].value
                        secretaria_s = working_sheet["J" + str(i)].value
                        vinculo_s = working_sheet["K" + str(i)].value
                        

                        fiscal_principal = Fiscal(principal, matricula_p, cargo_p, vinculo_p, secretaria_p)
                        fiscal_suplente = Fiscal(suplente, matricula_s, cargo_s, vinculo_s, secretaria_s)
                        
                        portaria.fiscais.append({
                            "principal": fiscal_principal, 
                            "suplente": fiscal_suplente
                        })
                                        
                else: 
                    fornecedores.append({
                        "fornecedor": working_sheet["B" + str(i)].value,
                        "cnpj": working_sheet["C" + str(i)].value,
                        "ata": working_sheet["D" + str(i)].value
                    })

            return portaria
                
        def criar_documento_portaria(portaria):
            portaria.criar_arquivo()
            pass 
        
        portaria = coletar_informacoes_portarias()
        criar_documento_portaria(portaria)

    
    PLANILHA = load_workbook(sheet_planilha, data_only= True)
    
    if (OPTION == "1"): gerenciador.entrarEmPasta("WORD") 
    
    if (OPTION == "1" or OPTION == "2"):
        for ordem in (PLANILHA.sheetnames):
            if (ordem.isdigit() and OPTION == "1"):
                gerar_termos(PLANILHA[ordem], ordem)
            elif (ordem.isdigit() and OPTION == "2"):
                buscar_informacoes_protocolo(PLANILHA[ordem])
                
    if (OPTION == "4"):
        print(OPTION, 'imsoconfused', PLANILHA, sheet_planilha)
        gerar_portarias(PLANILHA)
    
    if (OPTION == "2"):
        return param["ARR"]
    
    """
    """


    """
    
    
    
    
    def instanciar_documento(sheet):
        return Documento(sheet[MAPEAMENTO['contratado']].value)
    
    def verificar_se_arquivo_existe(contratado, sheet):       
        numero_af = sheet[MAPEAMENTO['numero_af']].value
        arq_prototipo = f'{nome}. {contratado} - AF {numero_af[:4]}.docx'      
        if (gerenciador.verificarArquivo(arq_prototipo)): return True   
        return arq_prototipo
        
    def colher_informacoes_relatorio(contratado, sheet):
        liq = sheet[MAPEAMENTO['numero_liquidacao']].value 
        val = sheet[MAPEAMENTO['valor_bruto_liquidacao']].value
        data = sheet[MAPEAMENTO['data_liquidacao']].value 
            
        rel = Relatorio(contratado, liq, val, data)
        return rel
        
    def colher_informacoes_termo(ordem, contratado, sheet, arq): 
        def customizar_mensagem(tipo):
            if (tipo == "Locação"): return "Por este instrumento, em caráter DEFINITIVO, atestamos que a locação acima identificada atende às exigências contratuais."
            return f"Por este instrumento, em caráter DEFINITIVO, atestamos que os {tipo.lower()} acima identificados atendem às exigências contratuais."
        endereco = rf"{gerenciador.pasta_atual}\{arq}.docx"
        numero_af = sheet[MAPEAMENTO['numero_af']].value
        contrato = sheet[MAPEAMENTO['n_contrato']].value 
        objeto = sheet[MAPEAMENTO['objeto']].value 
        mensagem = customizar_mensagem(str(sheet[MAPEAMENTO['tipo_nota']].value))
        return Termo(contratado, endereco, ordem, contrato, objeto, numero_af, mensagem, gerenciador.tipo_arquivo)
    
    for ordem in PLANILHA.sheetnames:
        try: 
            nome = int (ordem)
        except: 
            pass 
        else: 
            pass
        
            
            RELATORIO_INFO.append(colher_informacoes_relatorio(doc.contratado, PLANILHA[ordem]))
            nome_arquivo = verificar_se_arquivo_existe(doc.contratado, PLANILHA[ordem])
            #nome_arquivo = "1. CENTRO AMERICA FROTAS LTDA - AF 1950"
            # Verifica se o arquivo existe, se não cria um termo na pasta atual (dia atual/ remessa x / word)
            if (nome_arquivo != True):
                
                termo = colher_informacoes_termo(ordem, doc.contratado, PLANILHA[ordem], nome_arquivo)
                
                doc = termo.criar_arquivo(modelo_termo)
                doc.save(termo.endereco)
                try:
                    termo.salvar_pdf()
                except: 
                    pass
               
                
                
                ****** end_modelo needs to be in class
                
            
                #termo.copiar_arquivo(end_modelo, end_novo)
                
                ******
            
                pass
    """
            
        

        