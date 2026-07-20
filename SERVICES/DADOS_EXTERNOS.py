from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import time 

import requests
from openpyxl import load_workbook

from ENV.environment import pegar_pasta_downloads
from SERVICES.PLANILHA import CONVERTER_PARA_XLSX
from MODULES.GerenciarArquivos import verificarArquivo

# Aqui os dados são buscados de fontes externas, seja através de requisição por API ou por acesso automatizado ao site. 
# Os dados (planilha bruta ou JSON) retornam para o serviço MAIN que redireciona para as planilhas, que irá substituir os dados nas planilhas conforme necessário.

URL = "http://bataguassums.biosnet.com.br:8079"

## 1. to do:

    # criar modulo para converter xls para xlsx
    # ler a planilha de liquidaçao e voltar com um json util
    
def COLETAR_DADOS_EXTERNOS(tipo):

    def get_liquidacoes():
        return {"url": rf"{URL}/transparencia/DespesasLiquidadas.aspx",
                "botao": rf"return ProcessaDados('lnkDespesasLiquidadas')",
                "planilha_link": "Portal Transp. Despesas Liquidadas.xls",
                "dados_para_planilha": {
                    "A": "Local",
                    "B": "Fundo",
                    "C": "Empenho",
                    "D": "Data",
                    "E": "Valor",
                    "F": "Favorecido"
                }
            }
    
    def get_servidores():
        return { "url": rf"{URL}/transparencia/VersaoJson/Pessoal/",
                "json": {
                    "ConectarExercicio": "2026",
                    "Listagem": "Servidores",
                    "Ano": "2026",
                    "Empresa": "1",
                    "MostraDadosConsolidado": "True",
                    "MesFinalPeriodo": "01",
                    },   
                "dados_para_planilha": {
                    "Nome": "NOME", 
                    "Matricula": "ID",
                    "Unidade": "DIVISAO",
                    "Cargo": "CARGO",
                    "Vinculo": "VINCULO",
                }
            }   
         
    def get_contratos(): 
        return { "url": rf"{URL}/transparencia/VersaoJson/LicitacoesEContratos/",
                "json": {
                    "ConectarExercicio": "2026",
                    "Listagem": "Contratos",
                    "Ano": "2026",
                    "Empresa": "1",
                    "MostraDadosConsolidado": "True",
                    "ContratosApenasPublicados": "False"
                    },   
                "dados_para_planilha": {
                    "Processo": "PROCLIC", 
                    "Ano": "ANO",
                    "Contratado": "FORNECEDOR",
                    "Modalidade": "LICIT",
                    "Número Modalidade": "NUMLICMOD",
                    "Contrato": "CONTRATONUM",
                    "Objeto": "OBJETO_COMPLETO",
                    "Vigência": "VIGENF",
                    "Fiscais": "RESPON"
                },
            }
    
    def get_empenhos():
        return  { "url": rf"{URL}/transparencia/VersaoJson/Despesas/", 
                "json": {
                "ConectarExercicio": "2026",
                "Listagem": "DespesasGerais",
                "DiaInicioPeriodo": "01",
                "MesInicialPeriodo": "01",
                "DiaFinalPeriodo": "31",
                "MesFinalPeriodo": "12",
                "Ano": "2026",
                "Empresa": "1",
                "MostrarFornecedor": "True",
                "MostraDadosConsolidado": "False",
                "UFParaFiltroCOVID": "",
                "MostrarCNPJFornecedor": "True",
                "ApenasIDEmpenho": "False",
                }, 
                "dados_para_planilha": {
                    "Fornecedor": "NOMEFOR",
                    "Codigo": "CODIGO",
                    "Tipo_Empenho": "TPEM",
                    "Data": "DATAE",
                    "Projeto_Atividade": "PROJETO_ATIVIDADE_NOME",
                    "Programa": "PROGRAMANOME",
                    "Ficha": "FICHA",
                    "Fonte": "FONTE_STN",
                }
            }
    
    """
    
    if tipo == "MODO_WEB":
        try: 
            driver = webdriver.Chrome()

            # entra na raiz do sistema
            driver.get("http://bataguassums.biosnet.com.br:8079/transparencia/")

            # aguarda a página carregar
            WebDriverWait(driver, 10).until(
                lambda d: d.execute_script("return document.readyState") == "complete"
            )

            # agora navega para despesas
            driver.execute_script(direcao)
            driver.get(link)
            
            WebDriverWait(driver, 10).until(
                lambda d: d.execute_script("return document.readyState") == "complete"
            )
            # procura o botão
            botao = WebDriverWait(driver, 25).until(
                EC.element_to_be_clickable((By.ID, "btnExportarXLS"))
            )

            botao.click()
            time.sleep(10)
        except:
            print(f"Não foi possível capturar a planilha referente a... {link}. Permanecendo com os dados anteriores.")
    """
    
    print(f"♡ Realizando busca de dados --> {tipo}")
    if (tipo == "liquidacoes"):
        return COLETAR_DADOS_EXTERNOS_XLSX(get_liquidacoes())
    elif (tipo == "contratos"):
        return COLETAR_DADOS_EXTERNOS_JSON(get_contratos())
    elif (tipo == "empenhos"):
        return COLETAR_DADOS_EXTERNOS_JSON(get_empenhos())
    elif (tipo == "servidores"):
        return COLETAR_DADOS_EXTERNOS_JSON(get_servidores())
   

# Ambos retornam JSON, o primeiro extrai dados da API do transparência, o segundo extrai os dados de uma planilha
def COLETAR_DADOS_EXTERNOS_JSON(param):
    
    session = requests.Session()
    
    def buscar_dados(url, json, dados_planilha):
        response = session.get(url, params=json)
        dados_extraidos = response.json()
        #A função abaixo possui a função de relacionar os dados com as células da planilha
        dados_tratados = realizar_tratativa_nos_dados(url, dados_planilha, dados_extraidos)
        return dados_tratados

        
    def realizar_tratativa_nos_dados(url, dados_planilha, dados_extraidos):            
       
        dados_tratados = []
        def formatar_numero_contrato(dict):
            if (int(dict['Ano']) >= 2025):
                parse = f"{dict['Contrato'][2:]}/{dict['Ano'][2:]}"
                dict['Contrato'] = parse   
        for index in range(len(dados_extraidos)):
            dict = {}
            for chave in dados_planilha:       
                INFO = dados_extraidos[index][dados_planilha[chave]]
                dict[chave] = INFO
            if ("contratos" in url.lower()):
                formatar_numero_contrato(dict)    
            dados_tratados.append(dict)
        return dados_tratados
    
    dados_tratados = buscar_dados(param["url"], param["json"], param["dados_para_planilha"])
    return dados_tratados

def COLETAR_DADOS_EXTERNOS_XLSX(param):
    
    pasta_downloads = pegar_pasta_downloads()
    planilha = param["planilha_link"]
    
    def baixar_planilha(link, botao):
        def aguardar_pagina_carregar(driver):
            WebDriverWait(driver, 10).until(
                lambda d: d.execute_script("return document.readyState") == "complete"
            )  
        
        driver = webdriver.Chrome()
        driver.get(rf"{URL}/transparencia")
        aguardar_pagina_carregar(driver)
        
        driver.execute_script(botao)
        driver.get(link)
        
        botao = WebDriverWait(driver, 10).until(EC.element_to_be_clickable((By.ID, "btnExportarXLS")))
        botao.click()
        time.sleep(10)
    
    def converter_dados(estrutura_planilha): 
        
        conversao_planilha = CONVERTER_PARA_XLSX(rf"{pasta_downloads}\{planilha}")
        planilha_transp = load_workbook(rf"{conversao_planilha}")
        SHEET = planilha_transp.active 
        
        dados_convertidos = []
        for i in range(2, SHEET.max_row):
            dict = {}
            for chave in estrutura_planilha:
                dict[estrutura_planilha[chave]] = SHEET[chave + str(i)].value         
            dados_convertidos.append(dict)
            
        return dados_convertidos
    
    if ((verificarArquivo(rf"{pasta_downloads}\{planilha}x") or verificarArquivo(rf"{pasta_downloads}\{planilha}"))):
        pass 
    else:
        try:    baixar_planilha(param["url"], param["botao"])
        except: pass 

    return converter_dados(param["dados_para_planilha"])


r"""



#dto





def salvar_planilha(dados, caminho):
    # temp 
    
    def limpar_planilha(sheet):
        for row in sheet.iter_rows():
            for cell in row:
                cell.value = None
                
    PLANILHA = load_workbook(caminho)
    SHEET = PLANILHA.active
    
    limpar_planilha(SHEET)
  
    HEAD = []
    for keys in dados[0]:
        HEAD.append(keys)
    SHEET.append(HEAD)
    
    for linha in dados: 
        SHEET.append([linha.get(coluna, "") for coluna in HEAD])
    print(dados)
    PLANILHA.save(caminho)


#salvar_planilha(dados_contrato, r"C:\Users\Usuario\Documents\MRL\1. ANÁLISE DE PAGAMENTOS\INFO\BASE\ANÁLISE FISCAL.xlsx")
        

    
def get_liq_despesas():
     return  {
        "parametros": {
           "ConectarExercicio": "2026",
            "Listagem": "EmpenhosDespesas_Liquidado_PorNumeroEmpenho",
            "intNumeroEmpenho": "NumeroEmpenho",
            "strTipoEmpenho": "TipoEmpenho",
            "DiaInicioPeriodo": "01",
            "MesInicialPeriodo": "01",
            "DiaFinalPeriodo": "31",
            "MesFinalPeriodo": "12",
            "Ano": "2026",
            "Empresa": "1",
            "IDButton": "lnkDespesasPor_NotaEmpenho",
            "MostrarFornecedor": "True",
            "MostraDadosConsolidado": "False",
        },
        "celula_excel": {
            "a": ""
        } 
    }

def get_json_contratos():
    pass 

def pegar_informacoes(url, json):
    response = session.get(url, params=json["parametros"])
    dados = response.json()
    
    dados_para_planilha = []
    dict = {}
    
    parametro_dados_planilha = deepcopy(json["celula_excel"])

    for index in range(len(dados)):
        dict = {}
        for chave in parametro_dados_planilha:
            dict[chave] = dados[index][chave]
        dados_para_planilha.append(dict)
 
    
    print(dados_para_planilha[1]) 
       
    for index2 in range(len(dados_para_planilha)):
        break
        #print(dados_para_planilha[index2])
    


empenhos = pegar_informacoes("http://bataguassums.biosnet.com.br:8079/transparencia/VersaoJson/Despesas/", get_liq_despesas())
"""
#pegar_informacoes()


